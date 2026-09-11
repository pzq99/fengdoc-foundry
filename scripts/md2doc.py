#!/usr/bin/env python3
"""Convert manuscript or Supporting Information Markdown to formatted DOCX."""

from __future__ import annotations

import argparse
from copy import deepcopy
import re
import shutil
import subprocess
import sys
import tempfile
from pathlib import Path

from docx import Document
from docx.text.run import Run
from docx.enum.section import WD_ORIENT, WD_SECTION_START
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Mm, Pt, RGBColor


ROOT = Path(__file__).resolve().parents[1]
MANUSCRIPT_REFERENCE = ROOT / "assets" / "reference.docx"
SUPPORTING_REFERENCE = ROOT / "assets" / "reference-supporting.docx"
FONT = "Times New Roman"
BLUE = RGBColor(0x00, 0x00, 0x99)
LIGHT_GRAY = "F2F2F2"


def set_run_font(run, *, size=None, color=None):
    run.font.name = FONT
    if size is not None:
        run.font.size = Pt(size)
    if color is not None:
        run.font.color.rgb = color
    rpr = run._r.get_or_add_rPr()
    rfonts = rpr.get_or_add_rFonts()
    for key in ("ascii", "hAnsi", "cs", "eastAsia"):
        rfonts.set(qn(f"w:{key}"), FONT)


def set_cell_shading(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), fill)


def set_cell_border(cell, **edges):
    tc_pr = cell._tc.get_or_add_tcPr()
    borders = tc_pr.first_child_found_in("w:tcBorders")
    if borders is None:
        borders = OxmlElement("w:tcBorders")
        tc_pr.append(borders)
    for edge, attrs in edges.items():
        tag = f"w:{edge}"
        element = borders.find(qn(tag))
        if element is None:
            element = OxmlElement(tag)
            borders.append(element)
        for key, value in attrs.items():
            element.set(qn(f"w:{key}"), str(value))


def sect_geometry(sect_pr):
    pg = sect_pr.find(qn("w:pgSz"))
    mar = sect_pr.find(qn("w:pgMar"))
    if pg is None or mar is None:
        return None, False
    page_w = int(pg.get(qn("w:w")))
    page_h = int(pg.get(qn("w:h")))
    text_w = page_w - int(mar.get(qn("w:left"))) - int(mar.get(qn("w:right")))
    return text_w, page_w > page_h


def table_section_geometry(doc):
    """Return (text width, landscape) for each top-level table, in document order.

    A sectPr inside a paragraph's pPr ends a section at that paragraph, so a
    table belongs to the next section end at or after its position. Elements
    are matched by child ordinal: lxml proxies are not identity-stable across
    separate iterations, so id()-keyed lookups are unsafe here.
    """
    body = doc.element.body
    children = list(body.iterchildren())
    ends = []
    for pos, el in enumerate(children):
        if el.tag == qn("w:p"):
            p_pr = el.find(qn("w:pPr"))
            if p_pr is not None:
                sp = p_pr.find(qn("w:sectPr"))
                if sp is not None:
                    ends.append((pos, *sect_geometry(sp)))
    final = body.find(qn("w:sectPr"))
    if final is not None:
        ends.append((len(children), *sect_geometry(final)))
    if not ends:
        return [(9185, False)] * len(doc.tables)
    result = []
    index = 0
    for pos, el in enumerate(children):
        if el.tag == qn("w:tbl"):
            _, width, landscape = ends[min(index, len(ends) - 1)]
            result.append((width, landscape))
        if index < len(ends) and ends[index][0] == pos:
            index += 1
    return result


def column_ratios(table):
    ncols = len(table.columns)
    cols = []
    for element in table._tbl.tblGrid:
        try:
            cols.append(int(element.get(qn("w:w")) or 0))
        except ValueError:
            cols.append(0)
    cols = [c for c in cols if c > 0]
    if len(cols) != ncols or sum(cols) == 0:
        return [1.0 / ncols] * ncols
    return [c / sum(cols) for c in cols]


def set_table_width_dxa(table, total, ratios):
    """Fix the table to an exact width in twips with proportional columns."""
    tbl_pr = table._tbl.tblPr
    tbl_w = tbl_pr.find(qn("w:tblW"))
    if tbl_w is None:
        tbl_w = OxmlElement("w:tblW")
        tbl_pr.insert(0, tbl_w)
    tbl_w.set(qn("w:w"), str(total))
    tbl_w.set(qn("w:type"), "dxa")
    layout = tbl_pr.find(qn("w:tblLayout"))
    if layout is None:
        layout = OxmlElement("w:tblLayout")
        tbl_pr.append(layout)
    layout.set(qn("w:type"), "fixed")
    grid = table._tbl.tblGrid
    for element in list(grid):
        grid.remove(element)
    for ratio in ratios:
        col = OxmlElement("w:gridCol")
        col.set(qn("w:w"), str(round(total * ratio)))
        grid.append(col)
    for row in table.rows:
        for cell, ratio in zip(row.cells, ratios):
            tc_pr = cell._tc.get_or_add_tcPr()
            tc_w = tc_pr.find(qn("w:tcW"))
            if tc_w is None:
                tc_w = OxmlElement("w:tcW")
                tc_pr.insert(0, tc_w)
            tc_w.set(qn("w:w"), str(round(total * ratio)))
            tc_w.set(qn("w:type"), "dxa")


TC_PR_ORDER = ("tcW", "gridSpan", "hMerge", "vMerge", "tcBorders", "shd", "noWrap",
               "tcMar", "textDirection", "tcFitText", "vAlign", "hideMark")


def reorder_cell_tc_pr(cell):
    """Reorder tcPr children into the OOXML schema sequence."""
    tc_pr = cell._tc.tcPr
    if tc_pr is None:
        return

    def sort_key(element):
        tag = element.tag.rsplit("}", 1)[-1]
        return TC_PR_ORDER.index(tag) if tag in TC_PR_ORDER else len(TC_PR_ORDER)

    for element in sorted(list(tc_pr), key=sort_key):
        tc_pr.append(element)


def style_manuscript_tables(doc):
    # Six-column model-comparison tables need enough room for unbreakable
    # headers/method names (for example, ``Conditioning`` and ``DiffLinker``).
    # These content-safe proportions prevent Word from splitting such tokens
    # mid-word while retaining a portrait manuscript page.
    six_col_pct = [1350, 1350, 1650, 1600, 1800, 1435]
    seven_col_pct = [924, 846, 692, 539, 691, 539, 768]
    for table, (total, _landscape) in zip(doc.tables, table_section_geometry(doc)):
        table.style = doc.styles["Normal Table"]
        table.alignment = WD_TABLE_ALIGNMENT.CENTER
        table.autofit = False
        cell_font_size = 11 if len(table.columns) >= 6 else 12
        if len(table.columns) == 6:
            ratios = [p / sum(six_col_pct) for p in six_col_pct]
        elif len(table.columns) == 7:
            ratios = [p / sum(seven_col_pct) for p in seven_col_pct]
        else:
            ratios = column_ratios(table)
        set_table_width_dxa(table, total, ratios)

        for ri, row in enumerate(table.rows):
            row.height = Pt(11.35)
            for ci, cell in enumerate(row.cells):
                cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
                tc_pr = cell._tc.get_or_add_tcPr()
                # Allow Word to wrap at semantic boundaries. ``w:noWrap`` can
                # force narrow scientific tables to split words mid-token.
                nowrap = tc_pr.find(qn("w:noWrap"))
                if nowrap is not None:
                    tc_pr.remove(nowrap)
                if ri > 0 and ri % 2 == 1:
                    set_cell_shading(cell, LIGHT_GRAY)
                for paragraph in cell.paragraphs:
                    paragraph.style = doc.styles["Normal"]
                    pf = paragraph.paragraph_format
                    pf.alignment = WD_ALIGN_PARAGRAPH.CENTER
                    pf.line_spacing = 1.0
                    pf.space_before = Pt(2.4 if ri == 0 else 6)
                    pf.space_after = Pt(2.4 if ri == 0 else 6)
                    for run in paragraph.runs:
                        set_run_font(run, size=cell_font_size)
                        if ri == 0:
                            run.font.bold = True
                if ri == 0:
                    set_cell_border(
                        cell,
                        top={"val": "single", "sz": "8", "space": "0", "color": "auto"},
                        bottom={"val": "single", "sz": "8", "space": "0", "color": "auto"},
                    )
                elif ri == len(table.rows) - 1:
                    set_cell_border(
                        cell,
                        bottom={"val": "single", "sz": "4", "space": "0", "color": "auto"},
                    )
                reorder_cell_tc_pr(cell)


def set_repeat_table_header(row):
    tr_pr = row._tr.get_or_add_trPr()
    tbl_header = tr_pr.find(qn("w:tblHeader"))
    if tbl_header is None:
        tbl_header = OxmlElement("w:tblHeader")
        tr_pr.append(tbl_header)
    tbl_header.set(qn("w:val"), "true")


def clear_cell_borders(cell):
    set_cell_border(
        cell,
        top={"val": "nil"},
        left={"val": "nil"},
        bottom={"val": "nil"},
        right={"val": "nil"},
        insideH={"val": "nil"},
        insideV={"val": "nil"},
    )


def style_supporting_tables(doc):
    five_col_pct = [645, 2301, 623, 624, 807]
    for table, (total, landscape) in zip(doc.tables, table_section_geometry(doc)):
        table.style = doc.styles["Normal Table"]
        table.alignment = WD_TABLE_ALIGNMENT.CENTER
        table.autofit = False
        if len(table.columns) == 5 and landscape:
            # Measured grid of the reference wide table; applies to landscape
            # five-column tables only, per the format specification.
            ratios = [p / sum(five_col_pct) for p in five_col_pct]
        else:
            ratios = column_ratios(table)
        set_table_width_dxa(table, total, ratios)

        if table.rows:
            set_repeat_table_header(table.rows[0])
        for ri, row in enumerate(table.rows):
            for ci, cell in enumerate(row.cells):
                cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
                clear_cell_borders(cell)
                set_cell_shading(cell, LIGHT_GRAY if ri == 0 or (ri > 0 and ri % 2 == 0) else "FFFFFF")
                for paragraph in cell.paragraphs:
                    paragraph.style = doc.styles["Normal"]
                    pf = paragraph.paragraph_format
                    pf.alignment = WD_ALIGN_PARAGRAPH.LEFT if ci < 2 else WD_ALIGN_PARAGRAPH.CENTER
                    pf.line_spacing = 1.0
                    pf.space_before = Pt(3)
                    pf.space_after = Pt(3)
                    for run in paragraph.runs:
                        set_run_font(run, size=12)
                        if ri == 0:
                            run.font.bold = True
                if ri == 0:
                    set_cell_border(
                        cell,
                        top={"val": "single", "sz": "8", "space": "0", "color": "000000"},
                        bottom={"val": "single", "sz": "8", "space": "0", "color": "000000"},
                    )
                elif ri == len(table.rows) - 1:
                    set_cell_border(
                        cell,
                        bottom={"val": "single", "sz": "4", "space": "0", "color": "000000"},
                    )
                else:
                    set_cell_border(
                        cell,
                        bottom={"val": "single", "sz": "2", "space": "0", "color": "D9D9D9"},
                    )
                reorder_cell_tc_pr(cell)


def page_number(paragraph):
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = paragraph.add_run()
    begin = OxmlElement("w:fldChar")
    begin.set(qn("w:fldCharType"), "begin")
    instr = OxmlElement("w:instrText")
    instr.set(qn("xml:space"), "preserve")
    instr.text = " PAGE "
    separate = OxmlElement("w:fldChar")
    separate.set(qn("w:fldCharType"), "separate")
    text = OxmlElement("w:t")
    text.text = "1"
    end = OxmlElement("w:fldChar")
    end.set(qn("w:fldCharType"), "end")
    for element in (begin, instr, separate, text, end):
        run._r.append(element)
    set_run_font(run, size=12)


def configure_section(section, orientation="portrait"):
    if orientation == "landscape":
        section.orientation = WD_ORIENT.LANDSCAPE
        section.page_width = Mm(279.4)
        section.page_height = Mm(215.9)
        section.top_margin = Mm(24)
        section.bottom_margin = Mm(24)
        section.left_margin = Mm(24)
        section.right_margin = Mm(24)
        section.header_distance = Mm(12.7)
        section.footer_distance = Mm(12.7)
    else:
        section.orientation = WD_ORIENT.PORTRAIT
        section.page_width = Mm(210)
        section.page_height = Mm(297)
        section.top_margin = Mm(23)
        section.bottom_margin = Mm(23)
        section.left_margin = Mm(24)
        section.right_margin = Mm(24)
        section.header_distance = Mm(15.01)
        section.footer_distance = Mm(14.23)
    section.gutter = Mm(0)
    sect_pr = section._sectPr
    grid = sect_pr.find(qn("w:docGrid"))
    if grid is None:
        grid = OxmlElement("w:docGrid")
        sect_pr.append(grid)
    grid.set(qn("w:linePitch"), "312")


def add_section_at_marker(doc, paragraph, orientation):
    new_section = doc.add_section(WD_SECTION_START.NEW_PAGE)
    break_paragraph = doc.paragraphs[-1]
    paragraph._p.addprevious(break_paragraph._p)
    paragraph._p.getparent().remove(paragraph._p)
    configure_section(new_section, orientation)


def replace_section_markers(doc):
    markers = {
        "{{SECTION_BREAK_LANDSCAPE}}": "landscape",
        "{{SECTION_BREAK_PORTRAIT}}": "portrait",
    }
    for paragraph in list(doc.paragraphs):
        orientation = markers.get(paragraph.text.strip())
        if orientation:
            add_section_at_marker(doc, paragraph, orientation)


def replace_page_break_markers_with_page_break_before(doc):
    """Move SI hard breaks onto the next content paragraph.

    A standalone page-break run can be pushed to the next page when the
    preceding caption exactly fills a page, producing an accidental blank
    page. ``page_break_before`` on the next paragraph is stable in both
    cases and also works when that paragraph contains only an image.
    """
    paragraphs = list(doc.paragraphs)
    for index, paragraph in enumerate(paragraphs):
        if paragraph.text.strip() != "{{PAGE_BREAK}}":
            continue
        for next_paragraph in paragraphs[index + 1 :]:
            if next_paragraph.text.strip() not in {
                "{{PAGE_BREAK}}",
                "{{SECTION_BREAK_LANDSCAPE}}",
                "{{SECTION_BREAK_PORTRAIT}}",
            }:
                next_paragraph.paragraph_format.page_break_before = True
                break
        paragraph._p.getparent().remove(paragraph._p)


def set_page_geometry(doc, document_type):
    for section in doc.sections:
        configure_section(section, "portrait")
    if document_type == "supporting":
        replace_section_markers(doc)

    for index, section in enumerate(doc.sections):
        section.header.is_linked_to_previous = index > 0
        section.footer.is_linked_to_previous = index > 0
    footer = doc.sections[0].footer
    footer.is_linked_to_previous = False
    paragraph = footer.paragraphs[0]
    paragraph.clear()
    page_number(paragraph)


def style_heading(paragraph, level, *, page_start=False):
    pf = paragraph.paragraph_format
    pf.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY if level == 1 else WD_ALIGN_PARAGRAPH.LEFT
    pf.line_spacing = 1.5
    pf.space_after = Pt(7.8)
    pf.keep_with_next = True
    text = paragraph.text.strip()
    if level == 1:
        if page_start or text in {"Abstract", "1. Introduction"}:
            before = 7.8
        elif re.match(r"^[23]\.\s", text) or text == "References":
            before = 31.2
        else:
            before = 15.6
        pf.space_before = Pt(before)
    else:
        pf.space_before = Pt(7.8)
    for run in paragraph.runs:
        set_run_font(run, size=14 if level == 1 else 12, color=BLUE)
        if level <= 2:
            run.font.bold = True
        elif run.bold is None:
            run.font.bold = False


def split_run(run, offset):
    if offset <= 0 or offset >= len(run.text):
        return None
    new_r = deepcopy(run._r)
    run.text = run.text[:offset]
    new_run = Run(new_r, run._parent)
    new_run.text = new_run.text[offset:]
    run._r.addnext(new_r)
    return new_run


def split_paragraph_at(paragraph, absolute_offset):
    cursor = 0
    for run in list(paragraph.runs):
        end = cursor + len(run.text)
        if cursor < absolute_offset < end:
            split_run(run, absolute_offset - cursor)
            return
        cursor = end


def color_caption_lead(paragraph):
    match = re.match(r"^(Table|Figure)\s+S?\d+\.", paragraph.text.strip())
    if not match:
        return
    label_end = len(match.group(0))
    title_end = paragraph.text.find(".", label_end)
    title_end = len(paragraph.text) if title_end < 0 else title_end + 1
    split_paragraph_at(paragraph, label_end)
    split_paragraph_at(paragraph, title_end)
    cursor = 0
    for run in paragraph.runs:
        end = cursor + len(run.text)
        if cursor < title_end:
            run.font.color.rgb = BLUE
        if cursor < label_end:
            run.font.bold = True
        cursor = end


def bold_prefix(paragraph, pattern):
    match = re.match(pattern, paragraph.text.strip())
    if not match:
        return
    end_offset = match.end()
    split_paragraph_at(paragraph, end_offset)
    cursor = 0
    for run in paragraph.runs:
        end = cursor + len(run.text)
        if cursor < end_offset:
            run.font.bold = True
        cursor = end


def style_supporting_images(doc):
    max_width = Mm(162)
    for shape in doc.inline_shapes:
        if shape.width > max_width:
            ratio = max_width / shape.width
            shape.width = max_width
            shape.height = int(shape.height * ratio)
    for paragraph in doc.paragraphs:
        if paragraph._p.xpath(".//w:drawing | .//w:pict"):
            paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
            paragraph.paragraph_format.space_before = Pt(7.8)
            paragraph.paragraph_format.space_after = Pt(7.8)


def color_regex_matches(paragraph, pattern, color):
    matches = list(re.finditer(pattern, paragraph.text))
    for match in reversed(matches):
        split_paragraph_at(paragraph, match.end())
        split_paragraph_at(paragraph, match.start())
    cursor = 0
    for run in paragraph.runs:
        end = cursor + len(run.text)
        if any(cursor >= match.start() and end <= match.end() for match in matches):
            run.font.color.rgb = color
        cursor = end


# Internal cross-references ("Section 2.3", "Figure 4a,c", "Table S1",
# "Supplementary Figure S7", "Figures 2a-4c") are set in bold.
REFERENCE_PATTERN = re.compile(
    r"(?:Supplementary\s+)?(?:Figures?|Tables?)\s+S?\d+[a-z]?"
    r"(?:\s*[\u2013\-,]\s*(?:S?\d+)?[a-z]*(?:,[a-z])*)*"
    r"|\bSections?\s+\d+(?:\.\d+)?(?:\s*[\u2013\-]\s*\d+(?:\.\d+)*)?"
)

# Terms set in italic: Latin loci and statistical letters (P, q, n)
# preceding a comparison operator.
ITALIC_PATTERN = re.compile(
    r"\bin\s+(?:silico|vitro|vivo)\b"
    r"|\b[Pqn](?=\s*[=<>>\u2264\u2265])"
)

# Compact inline enumerators and narrative transition labels. Roman lists use
# the dedicated Unicode code points (ⅰ–ⅹ), which keeps them distinct from
# alphabetic panel labels such as Figure 2i.
ASCII_ROMAN_TO_UNICODE = {
    "i": "ⅰ", "ii": "ⅱ", "iii": "ⅲ", "iv": "ⅳ", "v": "ⅴ",
    "vi": "ⅵ", "vii": "ⅶ", "viii": "ⅷ", "ix": "ⅸ", "x": "ⅹ",
}
ASCII_ROMAN_TOKEN = r"(?:viii|vii|iii|vi|iv|ii|ix|v|i|x)"
ASCII_ROMAN_MARKER_PATTERN = re.compile(
    rf"\((?:{ASCII_ROMAN_TOKEN})(?:\s*(?:,|[–-])\s*{ASCII_ROMAN_TOKEN})*\)"
)
UNICODE_ROMAN_MARKER_PATTERN = re.compile(
    r"\((?:[ⅰ-ⅹ])(?:\s*(?:,|[–-])\s*[ⅰ-ⅹ])*\)"
)
LETTER_MARKER_PATTERN = re.compile(
    r"\((?:[a-i])(?:\s*(?:,|[–-])\s*[a-i])*\)"
)
ORDERED_TRANSITION_CANDIDATE = re.compile(r"\b(?:First|Second|Third)(?=\s*,)")

# URLs and e-mail addresses: dark blue, underlined, no hyperlink field.
# A trailing lookbehind keeps sentence punctuation outside the styled span.
LINK_PATTERN = re.compile(
    r"https?://[^\s)\]>]*(?<![.,;:])"
    r"|[A-Za-z0-9._%+-]+@[A-Za-z0-9.-]+\.[A-Za-z]{2,}"
)


def format_matches(doc, pattern, *, bold=False, italic=False, skip_styles=()):
    """Apply character formatting to complete regex matches."""
    for paragraph in doc.paragraphs:
        if paragraph.style is not None and paragraph.style.name in skip_styles:
            continue
        text = paragraph.text
        matches = list(pattern.finditer(text))
        if not matches:
            continue
        for match in reversed(matches):
            split_paragraph_at(paragraph, match.end())
            split_paragraph_at(paragraph, match.start())
        cursor = 0
        for run in paragraph.runs:
            end = cursor + len(run.text)
            if any(cursor >= match.start() and end <= match.end() for match in matches):
                if bold:
                    run.font.bold = True
                if italic:
                    run.font.italic = True
            cursor = end


def _sequence_markers(text):
    """Return markers only when their paragraph establishes an ordered list."""
    selected = []
    for pattern, token_pattern in (
        (UNICODE_ROMAN_MARKER_PATTERN, re.compile(r"[ⅰ-ⅹ]")),
        (LETTER_MARKER_PATTERN, re.compile(r"[a-i]")),
    ):
        markers = list(pattern.finditer(text))
        tokens = [m.group(0) for marker in markers for m in token_pattern.finditer(marker.group(0))]
        # A lone parenthetical character may be prose or a variable. Require
        # at least two distinct ordered items, whether separate, grouped, or ranged.
        if len(set(tokens)) >= 2:
            selected.extend((marker, token_pattern) for marker in markers)
    return selected


def format_enumerator_contents(doc):
    """Emphasize contextual enumerators, leaving their punctuation regular."""
    for paragraph in doc.paragraphs:
        token_spans = []
        punctuation_spans = []
        for marker, token_pattern in _sequence_markers(paragraph.text):
            marker_tokens = list(token_pattern.finditer(marker.group(0)))
            token_spans.extend(
                (marker.start() + token.start(), marker.start() + token.end())
                for token in marker_tokens
            )
            cursor = marker.start()
            for token in marker_tokens:
                start = marker.start() + token.start()
                end = marker.start() + token.end()
                if cursor < start:
                    punctuation_spans.append((cursor, start))
                cursor = end
            if cursor < marker.end():
                punctuation_spans.append((cursor, marker.end()))
        if not token_spans:
            continue
        boundaries = sorted({point for span in token_spans + punctuation_spans for point in span}, reverse=True)
        for point in boundaries:
            split_paragraph_at(paragraph, point)
        cursor = 0
        for run in paragraph.runs:
            end = cursor + len(run.text)
            if any(cursor >= start and end <= stop for start, stop in token_spans):
                run.font.bold = True
                run.font.italic = True
            elif any(cursor >= start and end <= stop for start, stop in punctuation_spans):
                run.font.bold = False
                run.font.italic = False
            cursor = end


def format_ordered_transitions(doc):
    """Bold First/Second/Third only when they form a sentence-led sequence."""
    for paragraph in doc.paragraphs:
        candidates = []
        for match in ORDERED_TRANSITION_CANDIDATE.finditer(paragraph.text):
            prefix = paragraph.text[:match.start()].rstrip()
            if not prefix or prefix[-1] in ".!?":
                candidates.append(match)
        labels = {match.group(0) for match in candidates}
        if "First" not in labels or len(labels) < 2:
            continue
        for match in reversed(candidates):
            split_paragraph_at(paragraph, match.end())
            split_paragraph_at(paragraph, match.start())
        cursor = 0
        for run in paragraph.runs:
            end = cursor + len(run.text)
            if any(cursor >= match.start() and end <= match.end() for match in candidates):
                run.font.bold = True
            cursor = end


def emphasize_references(doc):
    for paragraph in doc.paragraphs:
        matches = list(REFERENCE_PATTERN.finditer(paragraph.text))
        if not matches:
            continue
        for match in reversed(matches):
            split_paragraph_at(paragraph, match.end())
            split_paragraph_at(paragraph, match.start())
        cursor = 0
        for run in paragraph.runs:
            end = cursor + len(run.text)
            if any(cursor >= match.start() and end <= match.end() for match in matches):
                run.font.bold = True
            cursor = end


def style_links(doc):
    for paragraph in doc.paragraphs:
        matches = list(LINK_PATTERN.finditer(paragraph.text))
        if not matches:
            continue
        for match in reversed(matches):
            split_paragraph_at(paragraph, match.end())
            split_paragraph_at(paragraph, match.start())
        cursor = 0
        for run in paragraph.runs:
            end = cursor + len(run.text)
            if any(cursor >= match.start() and end <= match.end() for match in matches):
                run.font.color.rgb = BLUE
                run.font.underline = True
            cursor = end


def apply_text_conventions(doc):
    emphasize_references(doc)
    format_matches(doc, ITALIC_PATTERN, italic=True)
    format_enumerator_contents(doc)
    format_ordered_transitions(doc)
    style_links(doc)


def normalize_equations(doc):
    for paragraph in doc.paragraphs:
        if paragraph._p.xpath(".//m:oMath | .//m:oMathPara"):
            pf = paragraph.paragraph_format
            pf.alignment = WD_ALIGN_PARAGRAPH.CENTER
            pf.space_before = Pt(7.8)
            pf.space_after = Pt(7.8)
            pf.line_spacing = 1.5


def postprocess_manuscript(docx_path):
    doc = Document(docx_path)
    set_page_geometry(doc, "manuscript")

    in_references = False
    table_note_seen = False
    after_manual_page_break = False
    for paragraph in doc.paragraphs:
        text = paragraph.text.strip()
        if text == "{{PAGE_BREAK}}":
            in_references = False
            paragraph.clear()
            paragraph.add_run().add_break(WD_BREAK.PAGE)
            paragraph.paragraph_format.space_before = Pt(0)
            paragraph.paragraph_format.space_after = Pt(0)
            after_manual_page_break = True
            continue

        style_name = paragraph.style.name if paragraph.style else ""
        if style_name.startswith("Heading "):
            try:
                level = int(style_name.split()[-1])
            except ValueError:
                level = 1
            is_major_forced_start = level == 1 and bool(re.match(r"^[23]\.\s", text))
            page_start = after_manual_page_break or is_major_forced_start
            style_heading(paragraph, level, page_start=page_start)
            # Backward-compatible fallback for schema inputs that omit the
            # explicit marker. Normalized manuscripts should place the marker
            # before Results and Methods so it lives at the prior page end.
            if is_major_forced_start and not after_manual_page_break:
                paragraph.paragraph_format.page_break_before = True
            after_manual_page_break = False
            if text == "References":
                in_references = True
            continue

        if text:
            after_manual_page_break = False

        if style_name in {"Table Caption", "Figure Caption"}:
            in_references = False
        if in_references and text and style_name not in {"Table Caption", "Table Note", "Figure Caption"}:
            try:
                paragraph.style = doc.styles["Reference Entry"]
            except KeyError:
                pass

        pf = paragraph.paragraph_format
        if style_name in {"Normal", "Body Text", "First Paragraph", "Compact"} and text:
            pf.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
            pf.space_before = Pt(7.8)
            pf.space_after = Pt(7.8)
            pf.line_spacing = 1.5
            for run in paragraph.runs:
                set_run_font(run, size=12)
        elif style_name == "Manuscript Title":
            for run in paragraph.runs:
                set_run_font(run, size=13)
                run.font.bold = True
        elif style_name == "Authors":
            for run in paragraph.runs:
                set_run_font(run, size=12)
        elif style_name == "Affiliation":
            for run in paragraph.runs:
                set_run_font(run, size=11)
        elif style_name in {"Equal Contribution", "Correspondence", "Table Caption", "Table Note", "Figure Caption"}:
            for run in paragraph.runs:
                set_run_font(run, size=12)

        if style_name == "Correspondence":
            pass  # links and e-mails are styled globally by apply_text_conventions

        if style_name == "Table Caption":
            color_caption_lead(paragraph)
            table_note_seen = False
        elif style_name == "Figure Caption":
            color_caption_lead(paragraph)
        elif style_name == "Table Note":
            if not table_note_seen:
                pf.space_before = Pt(12)
                table_note_seen = True
            else:
                pf.space_before = Pt(0)
            pf.space_after = Pt(0)

    normalize_equations(doc)
    style_manuscript_tables(doc)
    apply_text_conventions(doc)
    doc.save(docx_path)


def postprocess_supporting(docx_path):
    doc = Document(docx_path)
    set_page_geometry(doc, "supporting")
    replace_page_break_markers_with_page_break_before(doc)

    table_note_seen = False
    for paragraph in doc.paragraphs:
        text = paragraph.text.strip()
        style_name = paragraph.style.name if paragraph.style else ""
        pf = paragraph.paragraph_format
        if style_name in {"Normal", "Body Text", "First Paragraph", "Compact"} and text:
            pf.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
            pf.space_before = Pt(7.8)
            pf.space_after = Pt(7.8)
            pf.line_spacing = 1.5
            for run in paragraph.runs:
                set_run_font(run, size=12)
        elif style_name == "SI Cover Label":
            pf.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
            pf.space_before = Pt(7.8)
            pf.space_after = Pt(23.4)
            pf.line_spacing = 1.2
            for run in paragraph.runs:
                set_run_font(run, size=12)
                run.font.bold = True
                run.font.italic = True
        elif style_name == "SI Cover Title":
            pf.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
            pf.space_before = Pt(7.8)
            pf.space_after = Pt(7.8)
            pf.line_spacing = 1.2
            for run in paragraph.runs:
                set_run_font(run, size=14)
                run.font.bold = True
        elif style_name == "SI Section Heading":
            pf.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
            pf.space_before = Pt(7.8)
            pf.space_after = Pt(15.6)
            pf.line_spacing = 1.5
            pf.keep_with_next = True
            for run in paragraph.runs:
                set_run_font(run, size=12, color=BLUE)
            bold_prefix(paragraph, r"^(?:Discussion|Method)\s+S\d+")
        elif style_name == "SI Subheading":
            pf.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
            pf.space_before = Pt(7.8)
            pf.space_after = Pt(7.8)
            pf.line_spacing = 1.5
            pf.keep_with_next = True
            for run in paragraph.runs:
                set_run_font(run, size=12)
        elif style_name == "SI Figure Caption":
            pf.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
            pf.space_before = Pt(7.8)
            pf.space_after = Pt(7.8)
            pf.line_spacing = 1.3208333
            for run in paragraph.runs:
                set_run_font(run, size=12)
            color_caption_lead(paragraph)
        elif style_name == "SI Table Caption":
            pf.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
            pf.space_before = Pt(6)
            pf.space_after = Pt(6)
            pf.line_spacing = 1.5
            pf.keep_with_next = True
            for run in paragraph.runs:
                set_run_font(run, size=12)
            color_caption_lead(paragraph)
            table_note_seen = False
        elif style_name == "SI Table Note":
            pf.alignment = WD_ALIGN_PARAGRAPH.LEFT
            if not table_note_seen:
                pf.space_before = Pt(12)
                table_note_seen = True
            else:
                pf.space_before = Pt(0)
            pf.space_after = Pt(0)
            pf.line_spacing = 1.5
            for run in paragraph.runs:
                set_run_font(run, size=12)

    normalize_equations(doc)
    style_supporting_images(doc)
    style_supporting_tables(doc)
    apply_text_conventions(doc)
    doc.save(docx_path)


def postprocess(docx_path, document_type):
    if document_type == "supporting":
        postprocess_supporting(docx_path)
    else:
        postprocess_manuscript(docx_path)


def normalize_ascii_roman_sequences(text):
    """Convert unambiguous ASCII Roman lists to dedicated Unicode numerals.

    A solitary ``(i)`` remains untouched because it may be an alphabetic panel
    label. Conversion requires a paragraph/block containing ``(i)`` plus at
    least one higher Roman item such as ``(ii)`` or ``(iv)``.
    """
    roman_values = {token: index for index, token in enumerate(ASCII_ROMAN_TO_UNICODE, start=1)}
    blocks = re.split(r"(\n\s*\n)", text)
    for index in range(0, len(blocks), 2):
        block = blocks[index]
        markers = list(ASCII_ROMAN_MARKER_PATTERN.finditer(block))
        tokens = [
            token
            for marker in markers
            for token in re.findall(ASCII_ROMAN_TOKEN, marker.group(0))
        ]
        values = {roman_values[token] for token in tokens}
        if 1 not in values or len(values) < 2:
            continue
        for marker in reversed(markers):
            replacement = re.sub(
                ASCII_ROMAN_TOKEN,
                lambda match: ASCII_ROMAN_TO_UNICODE[match.group(0)],
                marker.group(0),
            )
            block = block[:marker.start()] + replacement + block[marker.end():]
        blocks[index] = block
    return "".join(blocks)


def preprocess_markdown(source, destination, document_type):
    text = source.read_text(encoding="utf-8")
    text = normalize_ascii_roman_sequences(text)
    if document_type == "manuscript":
        # Keep SI inventories in the editable Markdown if useful, but exclude
        # them from the main-manuscript DOCX by default.
        text = re.split(
            r"(?m)^#\s+Supplementary (?:Figure Legends|Tables)\s*$",
            text,
            maxsplit=1,
        )[0].rstrip() + "\n"
    # Keep citation text but remove internal EndNote hyperlinks/bookmarks.
    # LibreOffice otherwise renders unresolved Word REF fields as a stray "X".
    text = re.sub(r"\[([^\]]+)\]\(#_ENREF_\d+\)", r"\1", text)
    text = re.sub(r'<span\s+id="_ENREF_\d+"\s+class="anchor"></span>', "", text)
    # The DOCX writer drops raw HTML formatting, so convert it to explicit
    # character styles defined in the reference document.
    for tag, style in (("sup", "Superscript"), ("sub", "Subscript"), ("u", "Underline")):
        pattern = re.compile(fr"<{tag}>(.*?)</{tag}>", re.DOTALL | re.IGNORECASE)
        text = pattern.sub(lambda m: f'[{m.group(1)}]{{custom-style="{style}"}}', text)
    # Pandoc's DOCX-to-Markdown output uses fenced `math` code blocks. Turn
    # those into display math so the DOCX writer emits native Word OMML.
    text = re.sub(
        r"```\s*math\s*\n(.*?)\n```",
        lambda m: "$$\n" + m.group(1) + "\n$$",
        text,
        flags=re.DOTALL,
    )
    # Normalize Pandoc's `$`...`$` inline-math representation.
    text = re.sub(r"\$`(.*?)`\$", lambda m: "$" + m.group(1) + "$", text, flags=re.DOTALL)
    destination.write_text(text, encoding="utf-8")


def main():
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("input", type=Path, help="Schema-compliant Markdown input")
    parser.add_argument("output", type=Path, help="Output DOCX path")
    parser.add_argument(
        "--type",
        dest="document_type",
        choices=("manuscript", "supporting"),
        default="manuscript",
        help="Document mode; defaults to manuscript for backward compatibility",
    )
    args = parser.parse_args()

    if not args.input.is_file():
        parser.error(f"input does not exist: {args.input}")
    reference = SUPPORTING_REFERENCE if args.document_type == "supporting" else MANUSCRIPT_REFERENCE
    build_script = "build_supporting_reference.py" if args.document_type == "supporting" else "build_reference.py"
    if not reference.is_file():
        parser.error(f"reference DOCX missing: {reference}; run scripts/{build_script}")
    pandoc = shutil.which("pandoc")
    if not pandoc:
        parser.error("pandoc is required on PATH")

    args.output.parent.mkdir(parents=True, exist_ok=True)
    with tempfile.TemporaryDirectory(prefix=f"md2doc-{args.document_type}-") as tmp:
        prepared = Path(tmp) / "prepared.md"
        preprocess_markdown(args.input, prepared, args.document_type)
        intermediate = Path(tmp) / "intermediate.docx"
        cmd = [
            pandoc,
            str(prepared),
            "--from=markdown+fenced_divs+bracketed_spans+tex_math_dollars",
            "--to=docx",
            f"--reference-doc={reference}",
            f"--resource-path={args.input.parent.resolve()}",
            "--wrap=none",
            "--output",
            str(intermediate),
        ]
        completed = subprocess.run(cmd, text=True, capture_output=True)
        if completed.returncode:
            sys.stderr.write(completed.stdout)
            sys.stderr.write(completed.stderr)
            raise SystemExit(completed.returncode)
        shutil.copy2(intermediate, args.output)
    postprocess(args.output, args.document_type)
    print(args.output.resolve())


if __name__ == "__main__":
    main()
